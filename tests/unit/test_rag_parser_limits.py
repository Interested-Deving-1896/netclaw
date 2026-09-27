"""Reject Office inputs exceeding configured limits before costly parsing."""
import sys
from pathlib import Path
import zipfile
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2]/'mcp-servers/rag-mcp'))
from ingestion import parsers


def test_zip_expansion_rejected_before_office_parser(tmp_path, monkeypatch):
    path = tmp_path/'expanded.docx'
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', b'x' * 200_000)
    assert path.stat().st_size < 10_000
    def forbidden(*args):
        pytest.fail('Oversized expanded archive reached Office parser')
    monkeypatch.setattr(parsers, '_parse_docx', forbidden)
    with pytest.raises(parsers.IngestError, match='SIZE_LIMIT_EXCEEDED'):
        parsers.parse_file(path, max_mb=10_000/(1024*1024))


def test_spreadsheet_sheet_limit(tmp_path):
    import openpyxl
    path = tmp_path/'sheets.xlsx'
    book = openpyxl.Workbook()
    book.active['A1'] = 'one'
    book.create_sheet('second')['A1'] = 'two'
    book.save(path)
    with pytest.raises(parsers.IngestError, match='SIZE_LIMIT_EXCEEDED'):
        parsers.parse_file(path, max_pages=1)


def test_slide_limit(tmp_path):
    from pptx import Presentation
    path = tmp_path/'slides.pptx'
    slides = Presentation()
    for i in range(2):
        slide = slides.slides.add_slide(slides.slide_layouts[0])
        slide.shapes.title.text = str(i)
    slides.save(path)
    with pytest.raises(parsers.IngestError, match='SIZE_LIMIT_EXCEEDED'):
        parsers.parse_file(path, max_pages=1)
